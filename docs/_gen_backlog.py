from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment, WrapText
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = "Product Backlog"

# 用户故事数据
stories = [
    {
        "编号": "US-01",
        "功能": "用户注册与登录",
        "角色": "系统管理员 / 驿站员工 / 普通用户",
        "故事": "作为系统管理员，我希望通过统一的登录入口登录系统，以便进入管理后台进行操作；作为普通用户，我希望能自行注册账号，以便查询本人快件状态。",
        "验收标准": "Given 已存在的账号 When 输入正确用户名和密码 Then 登录成功并返回 JWT 令牌\nGiven 已存在的账号 When 输入错误密码 Then 登录失败并提示具体原因\nGiven 新用户注册 When 用户名已存在 Then 注册失败并提示用户名已占用\nGiven 任意注册/登录请求 Then 密码均以 BCrypt 加盐哈希存储，数据库不存明文"
    },
    {
        "编号": "US-02",
        "功能": "角色权限管理",
        "角色": "系统管理员",
        "故事": "作为系统管理员，我希望为不同角色分配菜单与按钮权限，以便不同岗位只能操作其授权范围内的功能；同时我希望系统通过拦截器与注解自动鉴权，以保证未登录和越权行为被阻断。",
        "验收标准": "Given 未登录用户 When 访问需要认证的接口 Then 返回 401\nGiven 已登录但权限不足的用户 When 访问越权接口 Then 返回 403\nGiven 管理员为某角色分配了权限 When 该角色用户登录 Then 仅能看到授权菜单\nGiven 管理员账号 When 执行删除或禁用自身操作 Then 被系统拒绝"
    },
    {
        "编号": "US-03",
        "功能": "基础数据管理",
        "角色": "系统管理员 / 驿站员工",
        "故事": "作为系统管理员，我希望能维护驿站、货位等基础数据，以便支撑各驿站独立运营；作为驿站员工，我希望能维护本站的货位信息，以便合理分配存储空间。",
        "验收标准": "Given 新增驿站 When 驿站编号重复 Then 创建失败\nGiven 新增货位 When 库位编号在同一驿站下重复 Then 创建失败\nGiven 货位已满 When 收件登记请求分配货位 Then 应返回无可用货位提示"
    },
    {
        "编号": "US-04",
        "功能": "收件登记",
        "角色": "驿站员工",
        "故事": "作为驿站员工，我希望录入快件时系统能按快递公司自动校验运单号格式、生成唯一取件码并自动分配可用货位，以便快速完成入库、杜绝错误运单号、保证货位数据准确。",
        "验收标准": "Given 顺丰/京东/中通等 8 家快递公司 When 输入符合各自正则的运单号 Then 校验通过\nGiven 输入不符合任何公司正则的运单号 When 使用通用兜底规则 Then 校验通过\nGiven 输入不合法运单号格式 Then 校验失败并提示\nGiven 新快件 When 自动生成 8 位取件码 Then 取件码唯一（查库去重+重试）\nGiven 收件登记成功 When 货位 used_count 增加 Then 事务与入库同事务提交，任一步失败整体回滚"
    },
    {
        "编号": "US-05",
        "功能": "取件核销（出库）",
        "角色": "驿站员工",
        "故事": "作为驿站员工，我希望凭取件码与运单号双重匹配核销出库，并由系统自动核验身份、释放货位、计算逾期保管费，以便快速准确地完成取件、保证库存口径统一。",
        "验收标准": "Given 快件状态为在库待取 When 员工输入正确取件码+运单号 Then 核销成功，状态变更为已取件，货位 used_count-1\nGiven 仅输入取件码或仅输入运单号 When 执行核销 Then 核销失败，提示需双重匹配\nGiven 已取件快件 When 再次核销 Then 失败，提示快件已核销\nGiven 超过免费保管天数的快件 When 核销时 Then 自动计算逾期保管费并展示"
    },
    {
        "编号": "US-06",
        "功能": "发件登记",
        "角色": "驿站员工",
        "故事": "作为驿站员工，我希望能登记寄件单、自动生成寄件单号，并管理寄件单状态流转（待揽收→已揽收→已发出/已取消），以便规范寄件受理流程、保证状态可追溯。",
        "验收标准": "Given 新寄件单 When 登记成功 Then 自动生成格式为 S+yyyyMMddHHmmss+2 位随机的寄件单号\nGiven 寄件单状态为待揽收 When 执行揽收操作 Then 状态变为已揽收\nGiven 寄件单已揽收 When 执行发件出库操作 Then 状态变为已发出，记录发出时间与操作人\nGiven 寄件单在待揽收状态 When 执行取消操作 Then 状态变为已取消"
    },
    {
        "编号": "US-07",
        "功能": "异常件处理",
        "角色": "驿站员工",
        "故事": "作为驿站员工，我希望能将破损、丢失、地址错误、拒收、长期未取等快件登记为异常件，并记录处理结果，以便实现异常闭环管理、减少驿站运营风险。",
        "验收标准": "Given 在库或派送中的快件 When 登记为异常件 Then 状态变更为 EXCEPTION，同时货位 used_count 保持不变（异常件仍占用货位）\nGiven 已登记的异常件 When 录入处理结果 Then 异常记录包含处理人、处理方式、处理时间\nGiven 异常件处理完毕 When 需要退回 Then 状态变更为 RETURNED，货位 used_count-1"
    },
    {
        "编号": "US-08",
        "功能": "取件通知",
        "角色": "驿站员工 / 系统",
        "故事": "作为驿站员工，我希望系统能自动向收件人发送到件通知、逾期催取通知，且同一快件当天不重复催取，以便高效通知取件、释放货位资源；同时通知失败不应阻断主业务流程。",
        "验收标准": "Given 收件登记成功 When 系统自动触发到件通知 Then 通知记录写入数据库，状态为已发送\nGiven 逾期件 When 执行批量催取 Then 同一快件当天仅发送一次催取通知\nGiven 通知发送接口异常 When 主业务流程执行 Then 主业务不被阻断，通知记录标记为失败可重发"
    },
    {
        "编号": "US-09",
        "功能": "逾期保管费",
        "角色": "驿站员工 / 普通用户",
        "故事": "作为驿站员工，我希望系统能按日自动计算逾期件的保管费，并提供台账分页查询，以便准确计费、减少纠纷；作为普通用户，我希望查询取件时能看到逾期费用明细，以便知情缴费。",
        "验收标准": "Given 超过免费保管天数（默认 3 天）的快件 When 每日零点定时任务执行 Then overdue_days+1、storage_fee 按日累加\nGiven 逾期件 When 查询台账 Then 逾期件单独分页显示，包含逾期天数与累计保管费\nGiven 未逾期快件 When 核销取件 Then 保管费为 0"
    },
    {
        "编号": "US-10",
        "功能": "数据统计与看板",
        "角色": "系统管理员 / 驿站员工",
        "故事": "作为系统管理员，我希望能通过首页看板和数据大屏查看驿站运营全貌（今日入库/取件/寄件、在库/逾期/异常、派送中/累计等指标，以及出入库趋势、快递公司分布、快件类型分布图表），以便快速掌握运营状态、辅助决策；作为驿站员工，我希望能看到本站的运营数据，以便做好日常运营。",
        "验收标准": "Given 管理员登录首页 When 看板加载 Then 展示 8 项指标卡 + 3 张图表（近 7 日出入库折线、快递公司分布饼图、快件类型分布饼图）+ 货位使用率\nGiven 员工登录 When 查看看板 Then 仅展示本站数据（数据权限隔离）\nGiven 数据大屏 When 打开 Then 支持 30 秒自动轮询刷新，滚动榜单实时更新"
    }
]

# 表头
headers = ["故事编号", "所属功能", "角色", "用户故事", "验收标准"]

# 样式
header_fill = PatternFill(fill_type="solid", fgColor="1F4E79")
header_font = Font(bold=True, color="FFFFFF", size=11)
zebra1 = PatternFill(fill_type="solid", fgColor="FFFFFF")
zebra2 = PatternFill(fill_type="solid", fgColor="F7F9FC")
thin_border = Border(
    left=Side(style="thin", color="D9DEE7"),
    right=Side(style="thin", color="D9DEE7"),
    top=Side(style="thin", color="D9DEE7"),
    bottom=Side(style="thin", color="D9DEE7")
)
wrap_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
center_align = Alignment(horizontal="center", vertical="top", wrap_text=True)

# 写表头
for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.border = thin_border
    cell.alignment = center_align

# 写数据
for i, s in enumerate(stories):
    row = i + 2
    fill = zebra1 if i % 2 == 0 else zebra2
    values = [s["编号"], s["功能"], s["角色"], s["故事"], s["验收标准"]]
    for col, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=v)
        cell.fill = fill
        cell.border = thin_border
        cell.alignment = center_align if col <= 3 else wrap_align

# 列宽
widths = [12, 18, 28, 60, 70]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

# 行高（表头）
ws.row_dimensions[1].height = 24
# 数据行高
for r in range(2, len(stories) + 2):
    ws.row_dimensions[r].height = 80

# 冻结首行
ws.freeze_panes = "A2"

out = r"d:\wuyuhang-delivery-system\docs\Product Backlog-用户故事与验收标准.xlsx"
wb.save(out)
print("生成成功:", out)
print("共", len(stories), "条用户故事")
